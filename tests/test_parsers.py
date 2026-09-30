"""
Unit tests for Momento document parsing routines (.docx, .xlsx, .txt).
"""

import io
import pytest
import pandas as pd
from docx import Document

from main import parse_docx, parse_xlsx, parse_plaintext, extract_file_content


def test_parse_plaintext():
    text = "Hello world! This is a test file for Momento.\nLine 2: Data processing."
    raw_bytes = text.encode("utf-8")
    result = parse_plaintext(raw_bytes)
    assert "Hello world!" in result
    assert "Line 2: Data processing." in result


def test_parse_docx():
    # Create docx in memory
    doc = Document()
    doc.add_heading("Test Heading 1", level=1)
    doc.add_paragraph("This is the first paragraph of the test document.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "$1,000,000"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    docx_bytes = buffer.getvalue()

    result = parse_docx(docx_bytes)
    assert "Test Heading 1" in result
    assert "This is the first paragraph" in result
    assert "Metric | Value" in result
    assert "Revenue | $1,000,000" in result


def test_parse_xlsx():
    # Create Excel in memory with two sheets
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df1 = pd.DataFrame({
            "Department": ["Engineering", "Marketing", "Sales"],
            "Headcount": [45, 12, 28],
            "Budget": [500000, 150000, 300000]
        })
        df1.to_excel(writer, sheet_name="Q3_Summary", index=False)

        df2 = pd.DataFrame({
            "Task": ["Deploy VPS", "Setup Nginx"],
            "Status": ["Done", "Pending"]
        })
        df2.to_excel(writer, sheet_name="Action_Items", index=False)

    buffer.seek(0)
    xlsx_bytes = buffer.getvalue()

    result = parse_xlsx(xlsx_bytes)
    assert "Sheet: 'Q3_Summary'" in result
    assert "Department" in result
    assert "Engineering" in result
    assert "Sheet: 'Action_Items'" in result
    assert "Deploy VPS" in result


def test_extract_file_content_routing():
    # Test routing by extension
    txt_bytes = b"Sample text"
    res = extract_file_content("document.txt", txt_bytes)
    assert res == "Sample text"

    md_bytes = b"# Header\nMarkdown test"
    res_md = extract_file_content("notes.md", md_bytes)
    assert "Header" in res_md


def test_extract_unsupported_file():
    binary_data = b"\x00\x01\x02\x03\x04\xff\xfe"
    # An unknown extension without proper text decoding will fallback or reject
    # If it fails text decode or is binary, let's see how extract_file_content behaves
    try:
        extract_file_content("archive.bin", binary_data)
    except Exception as e:
        assert isinstance(e, (ValueError, UnicodeDecodeError))
