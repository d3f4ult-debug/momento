"""
Unit and integration tests for File Export & Download features in Momento.
"""

import os
import io
from unittest.mock import patch
import pytest
from fastapi.testclient import TestClient
from docx import Document
import openpyxl

from main import app
from services.file_exporter import (
    compile_docx,
    compile_xlsx,
    compile_csv,
    compile_text,
    detect_export_format,
    generate_export_file,
    parse_markdown_tables,
    ensure_downloads_dir,
    DOWNLOADS_DIR,
)

client = TestClient(app)


def test_detect_export_format():
    # Prompt-based detection
    assert detect_export_format("Please convert this to an excel spreadsheet") == "xlsx"
    assert detect_export_format("Format as word document (.docx)") == "docx"
    assert detect_export_format("Export data to csv") == "csv"
    assert detect_export_format("Save as markdown file") == "md"
    assert detect_export_format("Output as plain text") == "txt"

    # Input filename-based detection
    assert detect_export_format("Analyze this document", input_filename="report.docx") == "docx"
    assert detect_export_format("Summarize rows", input_filename="sales.xlsx") == "xlsx"
    assert detect_export_format("Process", input_filename="records.csv") == "csv"
    assert detect_export_format("Notes", input_filename="readme.md") == "md"

    # Content inference (tables -> xlsx)
    table_content = "| Col A | Col B |\n|---|---|\n| 1 | 2 |"
    assert detect_export_format("Summarize", content=table_content) == "xlsx"


def test_compile_docx(tmp_path):
    output_docx = str(tmp_path / "test_doc.docx")
    sample_text = (
        "# Executive Summary\n\n"
        "This is an introductory paragraph.\n\n"
        "## Key Findings\n"
        "- Finding 1: Revenue grew by 20%\n"
        "- Finding 2: Cost decreased by 5%\n\n"
        "### Metric Table\n"
        "| Metric | 2025 | 2026 |\n"
        "|---|---|---|\n"
        "| Revenue | $100M | $120M |\n"
        "| Margin | 15% | 18% |\n"
    )

    compile_docx(sample_text, output_docx, title="Test Report")
    assert os.path.exists(output_docx)
    assert os.path.getsize(output_docx) > 0

    # Verify docx content with python-docx
    doc = Document(output_docx)
    all_text = " ".join([p.text for p in doc.paragraphs])
    assert "Test Report" in all_text
    assert "Executive Summary" in all_text
    assert "Finding 1: Revenue grew by 20%" in all_text
    assert len(doc.tables) >= 1
    # Check table cells
    cell_values = [cell.text for row in doc.tables[0].rows for cell in row.cells]
    assert "Revenue" in cell_values
    assert "$120M" in cell_values


def test_compile_xlsx(tmp_path):
    output_xlsx = str(tmp_path / "test_sheet.xlsx")
    sample_markdown_table = (
        "| Quarter | Revenue | Profit |\n"
        "|---|---|---|\n"
        "| Q1 | 50000 | 12000 |\n"
        "| Q2 | 62000 | 15500 |\n"
        "| Q3 | 71000 | 19000 |\n"
    )

    compile_xlsx(sample_markdown_table, output_xlsx, sheet_title="Q_Financials")
    assert os.path.exists(output_xlsx)
    assert os.path.getsize(output_xlsx) > 0

    # Verify sheet with openpyxl
    wb = openpyxl.load_workbook(output_xlsx)
    assert "Q_Financials" in wb.sheetnames
    ws = wb["Q_Financials"]
    assert ws.cell(row=1, column=1).value == "Quarter"
    assert ws.cell(row=1, column=2).value == "Revenue"
    assert ws.cell(row=2, column=1).value == "Q1"
    assert str(ws.cell(row=2, column=2).value) == "50000"


def test_compile_csv(tmp_path):
    output_csv = str(tmp_path / "test_data.csv")
    sample_text = "Name,Role,City\nAlice,Engineer,Tashkent\nBob,Designer,Samarkand"
    compile_csv(sample_text, output_csv)

    assert os.path.exists(output_csv)
    with open(output_csv, "r", encoding="utf-8-sig") as f:
        content = f.read()
    assert "Alice" in content
    assert "Tashkent" in content


def test_generate_export_file_docx():
    result_text = "## Action Plan\n1. Launch MVP\n2. Monitor logs"
    export_data = generate_export_file(
        result_text=result_text,
        input_filename="project_spec.docx",
        prompt="Format as clean doc"
    )

    assert "file_id" in export_data
    assert export_data["format"] == "DOCX"
    assert export_data["filename"] == "transformed_project_spec.docx"
    assert export_data["file_id"] == "transformed_project_spec.docx"
    assert not any(c in export_data["filename"] for c in ["_", " "]) or not export_data["filename"].startswith("0")
    assert export_data["size_bytes"] > 0
    assert "download_url" in export_data
    assert export_data["download_url"] == "/api/download/transformed_project_spec.docx"

    # Clean up created file
    full_path = os.path.join(DOWNLOADS_DIR, export_data["file_id"])
    if os.path.exists(full_path):
        os.remove(full_path)


def test_strict_filename_integrity_translation():
    """Verify that translation preserves original filename with clean 'translated_' prefix and 0 hash clutter."""
    result_text = "Translated document body text."
    export_data = generate_export_file(
        result_text=result_text,
        input_filename="Quarterly_Report_2026.docx",
        prompt="Please translate this document to French"
    )

    assert export_data["filename"] == "translated_Quarterly_Report_2026.docx"
    assert export_data["file_id"] == "translated_Quarterly_Report_2026.docx"
    assert export_data["download_url"] == "/api/download/translated_Quarterly_Report_2026.docx"

    # Verify download endpoint serves clean filename
    dl_resp = client.get(export_data["download_url"])
    assert dl_resp.status_code == 200
    assert "translated_Quarterly_Report_2026.docx" in dl_resp.headers.get("content-disposition", "")

    full_path = os.path.join(DOWNLOADS_DIR, export_data["file_id"])
    if os.path.exists(full_path):
        os.remove(full_path)


def test_generate_export_file_xlsx():
    result_text = "| Product | Sales |\n|---|---|\n| Widget A | 1500 |\n| Widget B | 2300 |"
    export_data = generate_export_file(
        result_text=result_text,
        input_filename="sales_q3.xlsx",
        prompt="Extract sales data table"
    )

    assert export_data["format"] == "XLSX"
    assert export_data["filename"].endswith(".xlsx")
    assert export_data["size_bytes"] > 0

    full_path = os.path.join(DOWNLOADS_DIR, export_data["file_id"])
    if os.path.exists(full_path):
        os.remove(full_path)


def test_download_endpoint_success():
    # Create a dummy file in DOWNLOADS_DIR
    ensure_downloads_dir()
    test_id = "test_12345_sample_result.docx"
    test_file_path = os.path.join(DOWNLOADS_DIR, test_id)
    with open(test_file_path, "w", encoding="utf-8") as f:
        f.write("Test content for file download")

    try:
        response = client.get(f"/api/download/{test_id}")
        assert response.status_code == 200
        assert b"Test content for file download" in response.content
        assert "sample_result.docx" in response.headers.get("content-disposition", "")
    finally:
        if os.path.exists(test_file_path):
            os.remove(test_file_path)


def test_download_endpoint_not_found():
    response = client.get("/api/download/non_existent_file_99999.docx")
    assert response.status_code == 404
    assert "not found" in response.json()["detail"].lower()


def test_download_endpoint_path_traversal_blocked():
    # Attempt to traverse up out of the downloads directory
    response = client.get("/api/download/../../main.py")
    # Should either return 404 or be caught by safe_path check
    assert response.status_code in [404, 400]


@patch("main.execute_gemini_transformation")
def test_process_returns_export_file(mock_gemini):
    mock_gemini.return_value = (
        "# Analysis Report\n\n"
        "| KPI | Target | Actual |\n"
        "|---|---|---|\n"
        "| Growth | 10% | 14% |\n"
    )

    data = {
        "prompt": "Analyze business KPIs and output spreadsheet",
        "custom_api_key": "mock_key_for_test",
        "model": "gemini-2.5-flash",
    }
    response = client.post("/api/process", data=data)
    assert response.status_code == 200
    res_data = response.json()

    assert "export_file" in res_data
    assert res_data["export_file"] is not None
    assert "download_url" in res_data["export_file"]
    assert res_data["export_file"]["format"] in ["XLSX", "DOCX"]
    assert res_data["export_file"]["size_bytes"] > 0

    # Verify download works via the returned download_url
    dl_url = res_data["export_file"]["download_url"]
    dl_resp = client.get(dl_url)
    assert dl_resp.status_code == 200
    assert len(dl_resp.content) > 0

    # Clean up file
    file_id = res_data["export_file"]["file_id"]
    full_path = os.path.join(DOWNLOADS_DIR, file_id)
    if os.path.exists(full_path):
        os.remove(full_path)


@patch("main.is_google_authenticated", return_value={"authenticated": True})
@patch("main.execute_gemini_transformation")
def test_process_text_only_query_no_export_file(mock_gemini, mock_is_auth):
    mock_gemini.return_value = "Here are the files found in your Google Drive: 1. Budget 2026.xlsx, 2. Roadmap.docx"

    data = {
        "prompt": "what sheets in google do I have",
        "custom_api_key": "mock_key_for_test",
        "model": "gemini-2.5-flash",
    }
    response = client.post("/api/process", data=data)
    assert response.status_code == 200
    res_data = response.json()

    assert "export_file" in res_data
    assert res_data["export_file"] is None

